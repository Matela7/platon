"""Source reads require a successful RAG round in the current invocation."""

import asyncio
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest.mock import patch

from langchain.agents import create_agent
from langchain_core.language_models.fake_chat_models import FakeMessagesListChatModel
from langchain_core.messages import AIMessage, ToolMessage
from langchain_core.tools import tool
from langgraph.checkpoint.memory import InMemorySaver

from platon.agent.agents.database_agent import DatabaseAgent, SearchBeforeReadGuard
from platon.agent.models.agent_config import AgentConfig
from platon.agent.tools.file_management import READ_LIMITS, create_file_management_tools


class ScriptedModel(FakeMessagesListChatModel):
    def bind_tools(self, tools, **kwargs):
        return self


def call(name, identifier):
    return {"name": name, "args": {}, "id": identifier, "type": "tool_call"}


class DatabaseReadTests(unittest.TestCase):
    def make_agent(self, responses, root, *, failed=False):
        reads = []

        @tool
        def search_collection() -> list:
            """Search passages."""
            return {"status": "error", "errors": []} if failed else []

        @tool
        def read_file() -> str:
            """Read targeted source."""
            reads.append(True)
            return "source text"

        agent = DatabaseAgent(
            ScriptedModel(responses=responses), SimpleNamespace(),
            config=AgentConfig(backend_root=Path(root)), prompt="Search before reading.",
            tools_list=[search_collection, read_file], checkpointer=InMemorySaver(),
        )
        return agent, reads

    def test_graph_blocks_same_batch_then_allows_later_round_and_resets_next_run(self):
        with TemporaryDirectory() as root:
            agent, reads = self.make_agent([
                AIMessage(content="", tool_calls=[call("search_collection", "s1"), call("read_file", "r1")]),
                AIMessage(content="", tool_calls=[call("read_file", "r2")]),
                AIMessage(content="done"),
                AIMessage(content="", tool_calls=[call("read_file", "r3")]),
                AIMessage(content="done"),
            ], root)
            config = {"configurable": {"thread_id": "read-test"}}
            result = agent.invoke("First task", config=config)
            messages = {m.tool_call_id: m for m in result["messages"] if isinstance(m, ToolMessage)}
            self.assertEqual(messages["r1"].status, "error")
            self.assertEqual(messages["r2"].content, "source text")
            result = agent.invoke("New task", config=config)
            latest = next(m for m in result["messages"] if isinstance(m, ToolMessage) and m.tool_call_id == "r3")
            self.assertEqual(latest.status, "error")
            self.assertEqual(reads, [True])

    def test_async_graph_search_then_read(self):
        async def run():
            with TemporaryDirectory() as root:
                agent, reads = self.make_agent([
                    AIMessage(content="", tool_calls=[call("search_collection", "s1"), call("read_file", "r1")]),
                    AIMessage(content="", tool_calls=[call("read_file", "r2")]),
                    AIMessage(content="done"),
                    AIMessage(content="", tool_calls=[call("read_file", "r3")]),
                    AIMessage(content="done"),
                ], root)
                graph = create_agent(
                    model=agent.model, tools=agent.tools_list,
                    middleware=[SearchBeforeReadGuard()], checkpointer=InMemorySaver(),
                )
                result = await graph.ainvoke(
                    {"messages": [("user", "Inspect")]},
                    config={"configurable": {"thread_id": "async-test"}},
                )
                self.assertEqual(reads, [True])
                self.assertIn("source text", str(result))
                first_read = next(m for m in result["messages"] if isinstance(m, ToolMessage) and m.tool_call_id == "r1")
                self.assertEqual(first_read.status, "error")
                result = await graph.ainvoke(
                    {"messages": [("user", "New task")]},
                    config={"configurable": {"thread_id": "async-test"}},
                )
                stale_read = next(m for m in result["messages"] if isinstance(m, ToolMessage) and m.tool_call_id == "r3")
                self.assertEqual(stale_read.status, "error")
                self.assertEqual(reads, [True])
        asyncio.run(run())

    def test_serialized_search_error_does_not_unlock_read(self):
        with TemporaryDirectory() as root:
            agent, reads = self.make_agent([
                AIMessage(content="", tool_calls=[call("search_collection", "s1")]),
                AIMessage(content="", tool_calls=[call("read_file", "r1")]),
                AIMessage(content="done"),
            ], root, failed=True)
            result = agent.invoke("Inspect", config={"configurable": {"thread_id": "error-test"}})
            self.assertEqual(reads, [])
            self.assertTrue(any(isinstance(m, ToolMessage) and m.name == "read_file" and m.status == "error"
                                for m in result["messages"]))

    def test_read_modes_text_pages_pdf_pages_and_path_restriction(self):
        with TemporaryDirectory() as root:
            documents = Path(root) / "documents"
            documents.mkdir()
            (documents / "text.txt").write_text("x" * 40_000 + "\fsecond\fthird", encoding="utf-8")
            (documents / "fake.pdf").touch()
            read = next(t for t in create_file_management_tools(documents) if t.name == "read_file")
            for mode, limit in READ_LIMITS.items():
                output = read.invoke({"file_path": "text.txt", "mode": mode})
                self.assertTrue(output.startswith("x" * limit))
                self.assertIn("truncated", output)
            self.assertEqual(read.invoke({"file_path": "/documents/text.txt", "start_page": 2, "end_page": 2}), "second")
            denied = read.invoke({"file_path": "../outside.txt"})
            self.assertIn('"status": "error"', denied)
            invalid = read.invoke({"file_path": "text.txt", "start_page": 3, "end_page": 2})
            self.assertIn('"status": "error"', invalid)
            pages = [SimpleNamespace(extract_text=lambda: "first"), SimpleNamespace(extract_text=lambda: "second")]
            with patch("pypdf.PdfReader", return_value=SimpleNamespace(pages=pages)):
                self.assertEqual(read.invoke({"file_path": "fake.pdf", "start_page": 2, "end_page": 2}), "[Page 2]\nsecond")
            from pypdf import PdfWriter
            writer = PdfWriter()
            writer.add_blank_page(width=200, height=200)
            writer.add_blank_page(width=200, height=200)
            writer.write(documents / "real.pdf")
            self.assertEqual(read.invoke({"file_path": "real.pdf", "start_page": 2}), "[Page 2]\n")
            from docx import Document
            document = Document()
            document.add_paragraph("DOCX paragraph")
            document.add_table(rows=1, cols=1).cell(0, 0).text = "table value"
            document.save(documents / "real.docx")
            self.assertEqual(read.invoke({"file_path": "real.docx"}), "DOCX paragraph\ntable value")


if __name__ == "__main__":
    unittest.main()
